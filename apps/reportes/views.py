from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, FileResponse
from django.db.models import Sum, Count, Q
from django.utils import timezone
from datetime import timedelta
from io import BytesIO
import json
import csv


def _aplicar_filtro_fecha(request):
    """Devuelve las fechas de filtro según los parámetros GET."""
    filtro = request.GET.get('filtro', '30dias')
    hoy = timezone.now().date()
    
    if filtro == 'hoy':
        return hoy, hoy, 'Hoy'
    elif filtro == '7dias':
        return hoy - timedelta(days=7), hoy, 'Últimos 7 días'
    elif filtro == '30dias':
        return hoy - timedelta(days=30), hoy, 'Últimos 30 días'
    elif filtro == 'mes':
        inicio = hoy.replace(day=1)
        return inicio, hoy, 'Este mes'
    elif filtro == 'ano':
        inicio = hoy.replace(month=1, day=1)
        return inicio, hoy, 'Este año'
    else:
        return hoy - timedelta(days=30), hoy, 'Últimos 30 días'


def _get_datos_reportes(fecha_desde, fecha_hasta):
    """Recopila todos los datos para el reporte."""
    from apps.catalogo.models import Producto, Categoria
    from apps.inventario.models import Existencia, Movimiento
    from apps.despacho.models import LineaPedido

    categorias = Categoria.objects.filter(activo=True)
    stock_por_categoria = []
    movimientos_por_categoria = []

    for cat in categorias:
        total_stock = Existencia.objects.filter(
            producto__categoria=cat
        ).aggregate(total=Sum('cantidad'))['total'] or 0

        total_productos = Producto.objects.filter(
            categoria=cat, activo=True
        ).count()

        entradas = Movimiento.objects.filter(
            producto__categoria=cat, tipo='ENTRADA',
            fecha__date__gte=fecha_desde, fecha__date__lte=fecha_hasta
        ).aggregate(total=Sum('cantidad'))['total'] or 0

        salidas = Movimiento.objects.filter(
            producto__categoria=cat, tipo='SALIDA',
            fecha__date__gte=fecha_desde, fecha__date__lte=fecha_hasta
        ).aggregate(total=Sum('cantidad'))['total'] or 0

        stock_por_categoria.append({
            'categoria': cat,
            'stock': total_stock,
            'productos': total_productos,
            'entradas': entradas,
            'salidas': salidas,
        })

        movimientos_por_categoria.append({
            'categoria': cat.nombre,
            'entradas': entradas,
            'salidas': salidas,
        })

    top_productos = LineaPedido.objects.filter(
        pedido__fecha_pedido__gte=fecha_desde,
        pedido__fecha_pedido__lte=fecha_hasta
    ).values(
        'producto__sku',
        'producto__nombre',
        'producto__categoria__nombre'
    ).annotate(
        total_vendido=Sum('cantidad')
    ).order_by('-total_vendido')[:10]

    total_productos = Producto.objects.filter(activo=True).count()
    total_stock = Existencia.objects.aggregate(total=Sum('cantidad'))['total'] or 0
    total_movimientos = Movimiento.objects.filter(
        fecha__date__gte=fecha_desde, fecha__date__lte=fecha_hasta
    ).count()

    valor_inventario = 0
    for exist in Existencia.objects.select_related('producto'):
        valor_inventario += exist.cantidad * float(exist.producto.precio_venta)

    return {
        'stock_por_categoria': stock_por_categoria,
        'movimientos_por_categoria': movimientos_por_categoria,
        'top_productos': top_productos,
        'total_productos': total_productos,
        'total_stock': total_stock,
        'total_movimientos': total_movimientos,
        'valor_inventario': valor_inventario,
    }


@login_required
def reportes_inicio(request):
    """Panel de reportes."""
    fecha_desde, fecha_hasta, etiqueta_filtro = _aplicar_filtro_fecha(request)
    datos = _get_datos_reportes(fecha_desde, fecha_hasta)

    from apps.catalogo.models import Categoria
    categorias = Categoria.objects.filter(activo=True)

    labels_categorias = [c.nombre for c in categorias]
    data_stock = []
    data_entradas = []
    data_salidas = []

    for item in datos['stock_por_categoria']:
        data_stock.append(item['stock'])
        data_entradas.append(item['entradas'])
        data_salidas.append(item['salidas'])

    context = {
        **datos,
        'etiqueta_filtro': etiqueta_filtro,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'filtro_actual': request.GET.get('filtro', '30dias'),
        'labels_categorias_json': json.dumps(labels_categorias),
        'data_stock_categorias_json': json.dumps(data_stock),
        'data_entradas_categorias_json': json.dumps(data_entradas),
        'data_salidas_categorias_json': json.dumps(data_salidas),
    }

    return render(request, 'reportes/inicio.html', context)


# ==========================================
# EXPORTAR PDF
# ==========================================
@login_required
def exportar_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm

    fecha_desde, fecha_hasta, etiqueta = _aplicar_filtro_fecha(request)
    datos = _get_datos_reportes(fecha_desde, fecha_hasta)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    story = []

    styles = getSampleStyleSheet()
    
    # Título
    titulo_style = ParagraphStyle(
        'Titulo',
        parent=styles['Heading1'],
        textColor=colors.HexColor('#FF6B35'),
        fontSize=24,
        alignment=1,
        spaceAfter=10,
    )
    subtitulo_style = ParagraphStyle(
        'Subtitulo',
        parent=styles['Normal'],
        textColor=colors.HexColor('#6B7280'),
        fontSize=10,
        alignment=1,
        spaceAfter=20,
    )
    seccion_style = ParagraphStyle(
        'Seccion',
        parent=styles['Heading2'],
        textColor=colors.HexColor('#1F2937'),
        fontSize=14,
        spaceBefore=15,
        spaceAfter=10,
    )

    story.append(Paragraph('C.TECH', titulo_style))
    story.append(Paragraph('Reporte de Análisis - ' + etiqueta, subtitulo_style))
    story.append(Paragraph(f'Del {fecha_desde} al {fecha_hasta}', subtitulo_style))

    # KPIs
    story.append(Paragraph('Resumen General', seccion_style))
    kpi_data = [
        ['Métrica', 'Valor'],
        ['Total Productos', str(datos['total_productos'])],
        ['Total Stock', f"{datos['total_stock']} unidades"],
        ['Movimientos', str(datos['total_movimientos'])],
        ['Valor Inventario', f"${datos['valor_inventario']:.2f}"],
    ]
    kpi_table = Table(kpi_data, colWidths=[8*cm, 8*cm])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FF6B35')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F9FAFB')),
    ]))
    story.append(kpi_table)

    # Stock por categoría
    story.append(Paragraph('Stock por Categoría', seccion_style))
    cat_data = [['Categoría', 'Productos', 'Stock', 'Entradas', 'Salidas']]
    for item in datos['stock_por_categoria']:
        cat_data.append([
            item['categoria'].nombre,
            str(item['productos']),
            str(item['stock']),
            str(item['entradas']),
            str(item['salidas']),
        ])
    cat_table = Table(cat_data, colWidths=[4*cm, 3*cm, 3*cm, 3*cm, 3*cm])
    cat_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FF6B35')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F9FAFB')),
    ]))
    story.append(cat_table)

    # Top productos
    if datos['top_productos']:
        story.append(Paragraph('Top Productos Más Vendidos', seccion_style))
        top_data = [['SKU', 'Producto', 'Vendido']]
        for p in datos['top_productos']:
            top_data.append([
                p['producto__sku'],
                p['producto__nombre'][:40],
                str(p['total_vendido']),
            ])
        top_table = Table(top_data, colWidths=[4*cm, 9*cm, 3*cm])
        top_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FF6B35')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (2, 0), (2, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ]))
        story.append(top_table)

    story.append(Spacer(1, 20))
    story.append(Paragraph(
        f'Generado el {timezone.now().strftime("%d/%m/%Y %H:%M")} · C.TECH - Servicio y Calidad',
        subtitulo_style
    ))

    doc.build(story)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="reporte_ctech_{fecha_desde}_{fecha_hasta}.pdf"'
    return response


# ==========================================
# EXPORTAR EXCEL
# ==========================================
@login_required
def exportar_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    fecha_desde, fecha_hasta, etiqueta = _aplicar_filtro_fecha(request)
    datos = _get_datos_reportes(fecha_desde, fecha_hasta)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte C.TECH'

    # Estilos
    header_fill = PatternFill(start_color='FF6B35', end_color='FF6B35', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=12)
    title_font = Font(bold=True, color='FF6B35', size=18)
    border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB'),
    )

    # Título
    ws['A1'] = 'C.TECH'
    ws['A1'].font = title_font
    ws['A2'] = f'Reporte - {etiqueta}'
    ws['A3'] = f'Del {fecha_desde} al {fecha_hasta}'

    # Resumen
    row = 5
    ws[f'A{row}'] = 'RESUMEN GENERAL'
    ws[f'A{row}'].font = header_font
    ws[f'A{row}'].fill = header_fill
    ws[f'B{row}'].fill = header_fill
    row += 1

    resumen = [
        ('Total Productos', datos['total_productos']),
        ('Total Stock', datos['total_stock']),
        ('Movimientos', datos['total_movimientos']),
        ('Valor Inventario', f"${datos['valor_inventario']:.2f}"),
    ]
    for label, valor in resumen:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = valor
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        row += 1

    row += 2

    # Stock por categoría
    ws[f'A{row}'] = 'STOCK POR CATEGORÍA'
    ws[f'A{row}'].font = header_font
    ws[f'A{row}'].fill = header_fill
    row += 1

    headers = ['Categoría', 'Productos', 'Stock', 'Entradas', 'Salidas']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
    row += 1

    for item in datos['stock_por_categoria']:
        ws.cell(row=row, column=1, value=item['categoria'].nombre).border = border
        ws.cell(row=row, column=2, value=item['productos']).border = border
        ws.cell(row=row, column=3, value=item['stock']).border = border
        ws.cell(row=row, column=4, value=item['entradas']).border = border
        ws.cell(row=row, column=5, value=item['salidas']).border = border
        row += 1

    row += 2

    # Top productos
    if datos['top_productos']:
        ws[f'A{row}'] = 'TOP PRODUCTOS MÁS VENDIDOS'
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill
        row += 1

        headers = ['SKU', 'Producto', 'Vendido']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        row += 1

        for p in datos['top_productos']:
            ws.cell(row=row, column=1, value=p['producto__sku']).border = border
            ws.cell(row=row, column=2, value=p['producto__nombre']).border = border
            ws.cell(row=row, column=3, value=p['total_vendido']).border = border
            row += 1

    # Ajustar ancho
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="reporte_ctech_{fecha_desde}_{fecha_hasta}.xlsx"'
    return response


# ==========================================
# EXPORTAR WORD
# ==========================================
@login_required
def exportar_word(request):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    fecha_desde, fecha_hasta, etiqueta = _aplicar_filtro_fecha(request)
    datos = _get_datos_reportes(fecha_desde, fecha_hasta)

    doc = Document()

    # Título
    titulo = doc.add_heading('C.TECH', 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in titulo.runs:
        run.font.color.rgb = RGBColor(0xFF, 0x6B, 0x35)

    subtitulo = doc.add_paragraph(f'Reporte de Análisis - {etiqueta}')
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    fecha_p = doc.add_paragraph(f'Del {fecha_desde} al {fecha_hasta}')
    fecha_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Resumen
    doc.add_heading('Resumen General', level=1)
    resumen = [
        ('Total Productos', datos['total_productos']),
        ('Total Stock', f"{datos['total_stock']} unidades"),
        ('Movimientos', datos['total_movimientos']),
        ('Valor Inventario', f"${datos['valor_inventario']:.2f}"),
    ]
    for label, valor in resumen:
        p = doc.add_paragraph()
        p.add_run(f'{label}: ').bold = True
        p.add_run(str(valor))

    # Stock por categoría
    doc.add_heading('Stock por Categoría', level=1)
    tabla = doc.add_table(rows=1, cols=5)
    tabla.style = 'Light Grid Accent 1'
    headers = tabla.rows[0].cells
    for i, h in enumerate(['Categoría', 'Productos', 'Stock', 'Entradas', 'Salidas']):
        headers[i].text = h

    for item in datos['stock_por_categoria']:
        row = tabla.add_row().cells
        row[0].text = item['categoria'].nombre
        row[1].text = str(item['productos'])
        row[2].text = str(item['stock'])
        row[3].text = str(item['entradas'])
        row[4].text = str(item['salidas'])

    # Top productos
    if datos['top_productos']:
        doc.add_heading('Top Productos Más Vendidos', level=1)
        tabla2 = doc.add_table(rows=1, cols=3)
        tabla2.style = 'Light Grid Accent 1'
        headers = tabla2.rows[0].cells
        for i, h in enumerate(['SKU', 'Producto', 'Vendido']):
            headers[i].text = h

        for p in datos['top_productos']:
            row = tabla2.add_row().cells
            row[0].text = p['producto__sku']
            row[1].text = p['producto__nombre']
            row[2].text = str(p['total_vendido'])

    # Pie
    doc.add_paragraph()
    pie = doc.add_paragraph(
        f'Generado el {timezone.now().strftime("%d/%m/%Y %H:%M")} · C.TECH - Servicio y Calidad'
    )
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="reporte_ctech_{fecha_desde}_{fecha_hasta}.docx"'
    return response